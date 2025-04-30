from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file
from flask_login import login_required, current_user
from extensions import db
from models import User, Code, ScanLog, LoginLog
from datetime import datetime, time
import os, io, qrcode, random, string
from openpyxl import load_workbook

admin_bp = Blueprint('admin', __name__, url_prefix='/admin')

@admin_bp.route('/')
@login_required
def admin_dashboard():
    if current_user.role != 'admin':
        return redirect(url_for('scan.scanner'))
    codes = Code.query.all()
    users = User.query.all()
    logs = ScanLog.query.order_by(ScanLog.timestamp.desc()).limit(10).all()
    today = datetime.combine(datetime.now().date(), time.min)
    return render_template('admin_dashboard.html', codes=codes, users=users, logs=logs, today=today)

@admin_bp.route('/codes')
@login_required
def admin_codes():
    if current_user.role != 'admin':
        return redirect(url_for('scan.scanner'))
    page = request.args.get('page', 1, type=int)
    per_page = 20
    query = Code.query
    status_filter = request.args.get('status')
    if status_filter:
        query = query.filter(Code.scanned == (status_filter == 'scanned'))
    pagination = query.paginate(page=page, per_page=per_page)
    codes = pagination.items
    return render_template('admin_codes.html', codes=codes, pagination=pagination, status_filter=status_filter)

@admin_bp.route('/update_code/<code_id>', methods=['POST'])
@login_required
def update_code(code_id):
    if current_user.role != 'admin':
        return jsonify({'error': '权限不足'}), 403
    code = Code.query.get_or_404(code_id)
    data = request.json
    code.scanned = data.get('scanned', code.scanned)
    db.session.commit()
    return jsonify({'success': True})

@admin_bp.route('/users')
@login_required
def admin_users():
    if current_user.role != 'admin':
        return redirect(url_for('scan.scanner'))
    users = User.query.all()
    return render_template('admin_users.html', users=users)

@admin_bp.route('/create_user', methods=['POST'])
@login_required
def create_user():
    if current_user.role != 'admin':
        return jsonify({'error': '权限不足'}), 403
    data = request.json
    if User.query.filter_by(username=data['username']).first():
        return jsonify({'error': '用户名已存在'}), 400
    user = User(username=data['username'], role=data['role'])
    user.set_password(data['password'])
    db.session.add(user)
    db.session.commit()
    return jsonify({'success': True, 'user': {'id': user.id, 'username': user.username}})

@admin_bp.route('/delete_user/<user_id>', methods=['POST'])
@login_required
def delete_user(user_id):
    if current_user.role != 'admin' or str(current_user.id) == user_id:
        return jsonify({'error': '不能删除自己'}), 403
    user = User.query.get_or_404(user_id)
    db.session.delete(user)
    db.session.commit()
    return jsonify({'success': True})

@admin_bp.route('/login_logs')
@login_required
def login_logs():
    if current_user.role != 'admin':
        return redirect(url_for('scan.scanner'))
    page = request.args.get('page', 1, type=int)
    per_page = 25
    query = LoginLog.query.order_by(LoginLog.timestamp.desc())
    username_filter = current_user.username
    if username_filter:
        query = query.filter(LoginLog.username.ilike(f"%{username_filter}%"))
    pagination = query.paginate(page=page, per_page=per_page)
    logs = pagination.items
    return render_template('admin_login_logs.html', logs=logs, pagination=pagination, username_filter=username_filter)

@admin_bp.route('/delete_code/<code_id>', methods=['POST'])
@login_required
def delete_code(code_id):
    if current_user.role != 'admin':
        return jsonify({'error': '权限不足'}), 403
    code = Code.query.get_or_404(code_id)
    db.session.delete(code)
    db.session.commit()
    return jsonify({'success': True})

@admin_bp.route('/qrcode/<code>')
@login_required
def admin_qrcode(code):
    if current_user.role != 'admin':
        return '', 403
    img = qrcode.make(code)
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return send_file(buf, mimetype='image/png')

@admin_bp.route('/import', methods=['POST'])
@login_required
def import_excel():
    if 'file' not in request.files:
        return jsonify({'error': '没有文件'}), 400
    file = request.files['file']
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'error': '文件格式错误'}), 400
    filepath = os.path.join(os.getenv('UPLOAD_FOLDER', 'uploads/'), file.filename)
    file.save(filepath)
    success_count = 0
    fail_count = 0
    try:
        wb = load_workbook(filepath)
        ws = wb.active
        for row in ws.iter_rows(min_row=2):
            code = row[0].value
            if not code:
                code = ''.join(random.choices(string.ascii_uppercase + string.digits, k=8))
                while Code.query.filter_by(code=code).first():
                    code = ''.join(random.choices(string.ascii_uppercase + string.digits, k=8))
            if not Code.query.filter_by(code=code).first():
                db.session.add(Code(
                    code=code,
                    name=row[1].value,
                    phone=row[2].value,
                    note=row[3].value
                ))
                success_count += 1
            else:
                fail_count += 1
        db.session.commit()
        return jsonify({'success': True, 'success_count': success_count, 'fail_count': fail_count})
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500
