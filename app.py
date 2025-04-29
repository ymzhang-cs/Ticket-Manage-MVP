from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_from_directory
from flask_login import LoginManager, current_user, login_user, logout_user, login_required
from werkzeug.security import generate_password_hash, check_password_hash
import os
from openpyxl import load_workbook
from datetime import datetime, time
import random
import string
import io
import qrcode
from flask import send_file

from models import db, User, Code, ScanLog, ProfileLog, LoginLog

app = Flask(__name__)
app.config.from_object('config.Config')

# 初始化扩展
db.init_app(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'

@login_manager.user_loader
def load_user(id):
    return User.query.get(int(id))

# 基础路由
@app.route('/')
@login_required
def index():
    if current_user.role == 'admin':
        return redirect(url_for('admin_dashboard'))
    return redirect(url_for('scanner'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        user = User.query.filter_by(username=username).first()
        
        # 记录登录尝试
        log = LoginLog(
            username=username,
            ip_address=request.remote_addr,
            user_agent=request.user_agent.string
        )
        
        # 登录验证逻辑
        if user and user.check_password(password):
            login_user(user)
            log.success = True
            db.session.add(log)
            db.session.commit()
            return redirect(url_for('index'))
        else:
            log.success = False
            db.session.add(log)
            db.session.commit()
            flash('用户名或密码错误')
    
    return render_template('login.html')

# 新增登录日志页面
@app.route('/admin/login_logs')
@login_required
def login_logs():
    if current_user.role != 'admin':
        return redirect(url_for('scanner'))
    
    page = request.args.get('page', 1, type=int)
    per_page = 25
    
    # 支持筛选条件
    query = LoginLog.query.order_by(LoginLog.timestamp.desc())
    
    username_filter = current_user.username
    if username_filter:
        query = query.filter(LoginLog.username.ilike(f"%{username_filter}%"))
        
    pagination = query.paginate(page=page, per_page=per_page)
    logs = pagination.items
    
    return render_template('admin_login_logs.html', 
                         logs=logs, 
                         pagination=pagination,
                         username_filter=username_filter)
@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('您已退出登录')
    return redirect(url_for('login'))

# 扫码页面
@app.route('/scanner')
@login_required
def scanner():
    return render_template('scanner.html')

# 管理仪表盘
@app.route('/admin')
@login_required
def admin_dashboard():
    if current_user.role != 'admin':
        return redirect(url_for('scanner'))
    codes = Code.query.all()
    users = User.query.all()
    logs = ScanLog.query.order_by(ScanLog.timestamp.desc()).limit(10).all()
    today = datetime.combine(datetime.now().date(), time.min)  # 变为datetime类型
    return render_template('admin_dashboard.html', codes=codes, users=users, logs=logs, today=today)

@app.route('/admin/codes')
@login_required
def admin_codes():
    if current_user.role != 'admin':
        return redirect(url_for('scanner'))
    
    page = request.args.get('page', 1, type=int)
    per_page = 20
    
    # 支持筛选
    query = Code.query
    status_filter = request.args.get('status')
    if status_filter:
        query = query.filter(Code.scanned == (status_filter == 'scanned'))
    
    pagination = query.paginate(page=page, per_page=per_page)
    codes = pagination.items
    
    return render_template('admin_codes.html', 
                         codes=codes, 
                         pagination=pagination,
                         status_filter=status_filter)

# 单个二维码状态更新
@app.route('/admin/update_code/<code_id>', methods=['POST'])
@login_required
def update_code(code_id):
    if current_user.role != 'admin':
        return jsonify({'error': '权限不足'}), 403
        
    code = Code.query.get_or_404(code_id)
    data = request.json
    code.scanned = data.get('scanned', code.scanned)
    db.session.commit()
    return jsonify({'success': True})

@app.route('/admin/users')
@login_required
def admin_users():
    if current_user.role != 'admin':
        return redirect(url_for('scanner'))
    users = User.query.all()
    return render_template('admin_users.html', users=users)

# 创建用户
@app.route('/admin/create_user', methods=['POST'])
@login_required
def create_user():
    if current_user.role != 'admin':
        return jsonify({'error': '权限不足'}), 403
        
    data = request.json
    if User.query.filter_by(username=data['username']).first():
        return jsonify({'error': '用户名已存在'}), 400
        
    user = User(
        username=data['username'],
        role=data['role']
    )
    user.set_password(data['password'])
    db.session.add(user)
    db.session.commit()
    return jsonify({'success': True, 'user': {'id': user.id, 'username': user.username}})

# 删除用户
@app.route('/admin/delete_user/<user_id>', methods=['POST'])
@login_required
def delete_user(user_id):
    if current_user.role != 'admin' or str(current_user.id) == user_id:
        return jsonify({'error': '不能删除自己'}), 403
        
    user = User.query.get_or_404(user_id)
    db.session.delete(user)
    db.session.commit()
    return jsonify({'success': True})

@app.route('/user/profile', methods=['GET', 'POST'])
@login_required
def user_profile():
    if request.method == 'POST':
        form_type = request.form.get('form_type')  # 区分表单类型
        
        try:
            if form_type == 'username_form':
                new_username = request.form.get('new_username')
                
                if not new_username or new_username.strip() == '':
                    flash('用户名不能为空')
                    return redirect(url_for('user_profile'))
                    
                if new_username == current_user.username:
                    flash('新用户名与当前相同')
                    return redirect(url_for('user_profile'))
                    
                if User.query.filter_by(username=new_username).first():
                    flash('该用户名已存在')
                    return redirect(url_for('user_profile'))
                
                # 记录日志
                log = ProfileLog(
                    user_id=current_user.id,
                    action_type='username_change',
                    old_value=current_user.username,
                    new_value=new_username,
                    ip_address=request.remote_addr,
                    user_agent=request.user_agent.string
                )
                
                current_user.username = new_username
                db.session.add(log)
                flash('用户名修改成功')

            elif form_type == 'password_form':
                old_password = request.form.get('old_password')
                new_password = request.form.get('new_password')
                
                if not old_password or not new_password:
                    flash('请填写新旧密码')
                    return redirect(url_for('user_profile'))
                    
                if not current_user.check_password(old_password):
                    flash('旧密码错误')
                    return redirect(url_for('user_profile'))
                    
                if len(new_password) < 6:
                    flash('密码长度至少6位')
                    return redirect(url_for('user_profile'))
                
                # 记录日志
                log = ProfileLog(
                    user_id=current_user.id,
                    action_type='password_change',
                    old_value='',  # 不存储明文密码
                    new_value='********',
                    ip_address=request.remote_addr,
                    user_agent=request.user_agent.string
                )
                
                current_user.set_password(new_password)
                db.session.add(log)
                flash('密码修改成功')
                
            else:
                flash('未知操作类型')
            
            db.session.commit()
            
        except Exception as e:
            db.session.rollback()
            flash('操作失败: ' + str(e))
        
        return redirect(url_for('user_profile'))

    logs = ProfileLog.query.filter_by(user_id=current_user.id).order_by(ProfileLog.timestamp.desc()).limit(10).all()
    return render_template('user_profile.html', logs=logs)

# API接口（保持原有逻辑）
@app.route('/register', methods=['POST'])
@login_required
def register():
    if current_user.role != 'admin':
        return jsonify({'error': '权限不足'}), 403
        
    data = request.json
    if User.query.filter_by(username=data['username']).first():
        return jsonify({'error': '用户名已存在'}), 400
        
    user = User(username=data['username'], role=data['role'])
    user.set_password(data['password'])
    db.session.add(user)
    db.session.commit()
    return jsonify({'success': True})

# 二维码相关API
@app.route('/scan', methods=['POST'])
@login_required
def scan():
    code_str = request.json['code']
    code_record = Code.query.filter_by(code=code_str).first()
    
    if not code_record:
        return jsonify({'error': '无效二维码'}), 404
        
    if code_record.scanned:
        return jsonify({'error': '该二维码已使用'})
        
    code_record.scanned = True
    log = ScanLog(code_id=code_record.id, checker_id=current_user.id)
    db.session.add(log)
    db.session.commit()
    
    return jsonify({
        'success': True,
        'name': code_record.name,
        'note': code_record.note
    })

@app.route('/import', methods=['POST'])
@login_required
def import_excel():
    if 'file' not in request.files:
        return jsonify({'error': '没有文件'}), 400
        
    file = request.files['file']
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'error': '文件格式错误'}), 400
        
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
    file.save(filepath)
    
    success_count = 0
    fail_count = 0

    
    try:
        wb = load_workbook(filepath)
        ws = wb.active
        
        for row in ws.iter_rows(min_row=2):  # 跳过标题行
            code = row[0].value
            if not code:
                code = ''.join(random.choices(string.ascii_uppercase + string.digits, k=8))
                while Code.query.filter_by(code=code).first():
                    code = ''.join(random.choices(string.ascii_uppercase + string.digits, k=8))
            if not Code.query.filter_by(code=code).first():
                db.session.add(Code(
                    code=code,
                    name=row[1].value,
                    phone=row[2].value,
                    note=row[3].value
                ))
                success_count += 1
            else:
                fail_count += 1
        db.session.commit()
        return jsonify({'success': True, 'success_count': success_count, 'fail_count': fail_count})
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500
    
@app.route('/admin/delete_code/<code_id>', methods=['POST'])
@login_required
def delete_code(code_id):
    if current_user.role != 'admin':
        return jsonify({'error': '权限不足'}), 403
    code = Code.query.get_or_404(code_id)
    db.session.delete(code)
    db.session.commit()
    return jsonify({'success': True})

@app.route('/admin/qrcode/<code>')
@login_required
def admin_qrcode(code):
    if current_user.role != 'admin':
        return '', 403
    # 你可以根据实际需求调整二维码内容
    img = qrcode.make(code)
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return send_file(buf, mimetype='image/png')

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=8080)